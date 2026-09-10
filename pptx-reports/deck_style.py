"""
Shared visual-style helpers for the security-summary PPTX generators.

Design intent (per user request): a bright / light color tone, organized the
same way the toolkit's own HTML reports are (see allinonevmw.ps1's embedded
CSS for the VMSA / audit-reporter / KISA HTML reports - the hex values below
are taken directly from those <style> blocks so the PPTX decks feel like the
same product family), shapes that are not oversized, and a dense/compact
slide layout suitable for a short C-level readout rather than a slide-per-
finding wall of text.

This module has no CLI of its own - it is imported by
generate_security_decks.py's three deck builders.
"""

from __future__ import annotations

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
import copy

# --------------------------------------------------------------------------
# Canvas
# --------------------------------------------------------------------------
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.45)
CONTENT_W = SLIDE_W - 2 * MARGIN

FONT = "Segoe UI"
FONT_KR = "Malgun Gothic"  # used only on the KISA deck, which mirrors a Korean HTML report


def rgb(hex_str: str) -> RGBColor:
    hex_str = hex_str.lstrip("#")
    return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))


# Palette - pulled from the HTML reports' :root {} tokens so the decks read
# as the same product family (bright backgrounds, colored accents only on
# small elements, never a full-slide saturated fill).
P = {
    "bg": "F8FAFC",
    "bg_alt": "F1F5F9",
    "card": "FFFFFF",
    "border": "E2E8F0",
    "text": "1E293B",
    "text2": "334155",
    "muted": "64748B",
    "navy": "1E3A8A",
    "navy_dark": "1E2A5E",
    "slate": "1E293B",
    "slate2": "334155",
    "accent": "2563EB",
    "white": "FFFFFF",
    # severity (VMSA HTML)
    "crit_bg": "FEE2E2", "crit_fg": "DC2626",
    "high_bg": "FFEDD5", "high_fg": "EA580C",
    "med_bg": "FEF3C7", "med_fg": "D97706",
    "low_bg": "DCFCE7", "low_fg": "16A34A",
    "default_bg": "E2E8F0", "default_fg": "475569",
    # pass/fail/info (audit-reporter HTML)
    "pass_bg": "DCFCE7", "pass_fg": "16A34A",
    "fail_bg": "FEE2E2", "fail_fg": "DC2626",
    "info_bg": "CFFAFE", "info_fg": "0891B2",
    # KISA HTML (github-style status colors)
    "k_pass": "1A7F37", "k_fail": "CF222E", "k_warn": "9A6700",
    "k_manual": "0969DA", "k_error": "8250DF",
    "k_grid": "B7C0CA",
}

SEVERITY_COLORS = {
    "critical": (P["crit_bg"], P["crit_fg"]),
    "important": (P["high_bg"], P["high_fg"]),
    "high": (P["high_bg"], P["high_fg"]),
    "moderate": (P["med_bg"], P["med_fg"]),
    "medium": (P["med_bg"], P["med_fg"]),
    "low": (P["low_bg"], P["low_fg"]),
}

KISA_STATUS_COLORS = {
    "PASS": P["k_pass"], "FAIL": P["k_fail"], "WARN": P["k_warn"],
    "MANUAL": P["k_manual"], "ERROR": P["k_error"],
}

PRIORITY_COLORS = {
    "P0": (P["crit_bg"], P["crit_fg"]),
    "P1": (P["high_bg"], P["high_fg"]),
    "P2": (P["med_bg"], P["med_fg"]),
    "ADVANCED": (P["default_bg"], P["default_fg"]),
}


# --------------------------------------------------------------------------
# Presentation / slide plumbing
# --------------------------------------------------------------------------
def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def add_blank_slide(prs: Presentation):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    set_fill(slide.background, P["bg"])
    return slide


def set_fill(fill_holder, hex_color, transparency=None):
    """fill_holder is anything with a .fill (shape/background)."""
    fill = fill_holder.fill
    fill.solid()
    fill.fore_color.rgb = rgb(hex_color)
    if transparency is not None:
        _set_transparency(fill.fore_color, transparency)
    return fill


def _set_transparency(color_format, pct):
    """pct: 0-100 (100 = fully transparent). python-pptx has no public API
    for this, so we poke the alpha value into the underlying XML directly."""
    alpha = str(int(round((100 - pct) * 1000)))
    srgb = color_format._xFill.find(qn("a:srgbClr"))
    if srgb is None:
        return
    for tag in ("a:alpha",):
        existing = srgb.find(qn(tag))
        if existing is not None:
            srgb.remove(existing)
    node = srgb.makeelement(qn("a:alpha"), {"val": alpha})
    srgb.append(node)


def no_line(shape):
    shape.line.fill.background()


def set_line(shape, hex_color, weight_pt=0.75):
    shape.line.color.rgb = rgb(hex_color)
    shape.line.width = Pt(weight_pt)


def add_rect(slide, x, y, w, h, fill_hex=None, line_hex=None, line_pt=0.75,
             rounded=False, radius=0.06, shadow=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, x, y, w, h)
    if rounded:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if fill_hex:
        set_fill(shp, fill_hex)
    else:
        shp.fill.background()
    if line_hex:
        set_line(shp, line_hex, line_pt)
    else:
        no_line(shp)
    shp.shadow.inherit = False
    return shp


def add_gradient_band(slide, x, y, w, h, hex1, hex2, angle=45):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    no_line(shp)
    shp.shadow.inherit = False
    fill = shp.fill
    fill.gradient()
    stops = fill.gradient_stops
    stops[0].color.rgb = rgb(hex1)
    stops[0].position = 0.0
    stops[1].color.rgb = rgb(hex2)
    stops[1].position = 1.0
    try:
        fill.gradient_angle = angle
    except Exception:
        pass
    return shp


def add_text(slide, x, y, w, h, text, size=12, color=P["text"], bold=False,
             italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             font=FONT, spacing=None, wrap=True, shrink=False, line_spacing=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font
        run.font.color.rgb = rgb(color) if isinstance(color, str) else color
        if spacing is not None:
            _set_letter_spacing(run, spacing)
    return tb


def _set_letter_spacing(run, pts_hundredths):
    rPr = run._r.get_or_add_rPr()
    rPr.set("spc", str(int(pts_hundredths)))


def add_rich_text(slide, x, y, w, h, runs_spec, size=11, align=PP_ALIGN.LEFT,
                   anchor=MSO_ANCHOR.TOP, font=FONT, wrap=True):
    """runs_spec: list of (text, color_hex, bold) tuples on a single line."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    for text, color, bold in runs_spec:
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = font
        run.font.color.rgb = rgb(color)
    return tb


# --------------------------------------------------------------------------
# Composite widgets
# --------------------------------------------------------------------------
def _fit_title_size(title, text_w_in):
    """Picks the largest font size (from a small ladder) that keeps `title`
    on a single line within text_w_in inches, using a rough chars-per-inch
    estimate for bold Segoe UI. Prevents the title from wrapping onto the
    subtitle's line (there is no reliable text-measurement API in
    python-pptx, so this is a deliberately conservative heuristic)."""
    for size, chars_per_inch in ((22, 5.3), (20, 5.8), (18, 6.4), (16, 7.1)):
        if len(title) <= chars_per_inch * text_w_in:
            return size
    return 15


def _fit_single_line(text, text_w_in, chars_per_inch):
    """Truncates `text` with an ellipsis if it would not fit on one line at
    the given chars-per-inch estimate, so it never wraps onto/overlaps a
    fixed-position line below it."""
    max_chars = int(chars_per_inch * text_w_in)
    if len(text) <= max_chars:
        return text
    return text[: max(0, max_chars - 1)].rstrip() + "…"


def add_header_band(slide, kicker, title, subtitle, meta_lines,
                     color1=P["navy"], color2=P["navy_dark"], height=Inches(1.15)):
    """Compact gradient header band (not a full-slide fill - keeps the deck
    bright) with a small translucent meta box on the right, mirroring the
    HTML reports' header/.hero treatment."""
    add_gradient_band(slide, 0, 0, SLIDE_W, height, color1, color2, angle=45)

    text_w = SLIDE_W - Inches(4.3) - MARGIN
    if kicker:
        add_text(slide, MARGIN, Inches(0.16), text_w, Inches(0.22), kicker,
                  size=10.5, color="BFD1F7", bold=True, spacing=60)
    text_w_in = text_w / Inches(1)
    title_size = _fit_title_size(title, text_w_in)
    add_text(slide, MARGIN, Inches(0.38), text_w, Inches(0.46), title,
              size=title_size, color=P["white"], bold=True, wrap=False)
    if subtitle:
        subtitle = _fit_single_line(subtitle, text_w_in, chars_per_inch=10.5)
        add_text(slide, MARGIN, Inches(0.84), text_w, Inches(0.28), subtitle,
                  size=11.5, color="D6E0F5", wrap=False)

    if meta_lines:
        box_w = Inches(3.85)
        box_h = Inches(0.86)
        box_x = SLIDE_W - MARGIN - box_w
        box_y = (height - box_h) / 2
        add_rect(slide, box_x, box_y, box_w, box_h, fill_hex=P["white"], rounded=True, radius=0.18)
        _set_shape_transparency_by_index(slide, transparency=30)
        pad = Inches(0.16)
        add_text(slide, box_x + pad, box_y + Inches(0.08), box_w - 2 * pad, box_h - Inches(0.16),
                  meta_lines, size=9.5, color=P["slate2"], line_spacing=1.15)


def _set_shape_transparency_by_index(slide, transparency):
    """Applies transparency to the most-recently-added shape (used right
    after add_rect for the header meta box, to mimic rgba(255,255,255,.7))."""
    shp = slide.shapes[-1]
    _set_transparency(shp.fill.fore_color, transparency)


def add_page_header(slide, title, subtitle=None, accent=P["navy"]):
    """Small, compact section header for content slides - a short colored
    tick + bold title + thin rule, NOT a full band, so most of the slide
    stays available for dense content."""
    y = Inches(0.32)
    add_rect(slide, MARGIN, y + Inches(0.03), Inches(0.09), Inches(0.30), fill_hex=accent)
    add_text(slide, MARGIN + Inches(0.22), y, Inches(9.5), Inches(0.34), title,
              size=17, color=P["text"], bold=True)
    if subtitle:
        add_text(slide, MARGIN + Inches(0.22), y + Inches(0.33), Inches(9.5), Inches(0.22),
                  subtitle, size=10.5, color=P["muted"])
    rule_y = y + Inches(0.62) if subtitle else y + Inches(0.40)
    add_rect(slide, MARGIN, rule_y, CONTENT_W, Pt(1.1), fill_hex=P["border"])
    return rule_y + Inches(0.14)


def add_footer(slide, left_text, page_no=None, total_pages=None):
    y = SLIDE_H - Inches(0.36)
    add_rect(slide, MARGIN, y, CONTENT_W, Pt(0.75), fill_hex=P["border"])
    add_text(slide, MARGIN, y + Inches(0.06), Inches(10.5), Inches(0.24), left_text,
              size=8.5, color=P["muted"])
    if page_no is not None:
        add_text(slide, SLIDE_W - MARGIN - Inches(1.2), y + Inches(0.06), Inches(1.2), Inches(0.24),
                  f"{page_no} / {total_pages}", size=8.5, color=P["muted"], align=PP_ALIGN.RIGHT)


def add_kpi_row(slide, x, y, w, h, cards):
    """cards: list of dicts {label, value, color(optional), sub(optional)}.
    Small flat cards with a colored top accent bar - mirrors the HTML
    .stat-card { border-top: 4px solid ... } pattern."""
    n = len(cards)
    gap = Inches(0.16)
    card_w = (w - gap * (n - 1)) / n
    for i, c in enumerate(cards):
        cx = x + i * (card_w + gap)
        color = c.get("color", P["accent"])
        add_rect(slide, cx, y, card_w, h, fill_hex=P["card"], line_hex=P["border"], line_pt=0.75, rounded=True, radius=0.09)
        add_rect(slide, cx + Inches(0.02), y, card_w - Inches(0.04), Inches(0.06), fill_hex=color)
        pad = Inches(0.14)
        add_text(slide, cx + pad, y + Inches(0.16), card_w - 2 * pad, Inches(0.20), c["label"].upper(),
                  size=8.5, color=P["muted"], bold=True, spacing=30)
        add_text(slide, cx + pad, y + Inches(0.36), card_w - 2 * pad, Inches(0.5), str(c["value"]),
                  size=25, color=color, bold=True)
        if c.get("sub"):
            add_text(slide, cx + pad, y + h - Inches(0.28), card_w - 2 * pad, Inches(0.22), c["sub"],
                      size=8.5, color=P["muted"])


def bullet_card_height(n_bullets, title_size=12, bullet_size=10, min_h=Inches(1.7), max_h=Inches(2.8)):
    """Sizes a bullet card to its content instead of a fixed tall box, so
    short 2-4 bullet cards don't leave a large empty gap at the bottom."""
    est = Inches(0.5) + n_bullets * Inches(bullet_size / 58.0 + 0.14)
    return max(min_h, min(max_h, est))


def add_progress_bar(slide, x, y, w, h, label, pct_value, color, value_text=None):
    """A slim horizontal 'gauge' bar (label left, track + fill, % right) -
    used to give cover slides a compact dashboard feel without duplicating
    the fuller donut/table breakdowns on the overview slide."""
    label_w = Inches(2.1)
    pct_w = Inches(0.9)
    track_x = x + label_w
    track_w = w - label_w - pct_w - Inches(0.15)
    add_text(slide, x, y, label_w - Inches(0.1), h, label, size=10.5, color=P["text2"], anchor=MSO_ANCHOR.MIDDLE)
    track_h = Inches(0.16)
    track_y = y + (h - track_h) / 2
    add_rect(slide, track_x, track_y, track_w, track_h, fill_hex=P["border"], rounded=True, radius=0.5)
    fill_w = max(Inches(0.05), Emu(int(track_w * max(0.0, min(100.0, pct_value)) / 100.0)))
    add_rect(slide, track_x, track_y, fill_w, track_h, fill_hex=color, rounded=True, radius=0.5)
    add_text(slide, track_x + track_w + Inches(0.12), y, pct_w, h, value_text or f"{pct_value:.0f}%",
              size=10.5, color=color, bold=True, anchor=MSO_ANCHOR.MIDDLE, wrap=False)


def add_insight_bar(slide, x, y, w, h, label, text, color=P["navy"], bg=None):
    """A full-width 'key takeaway' banner - a light tinted card with a
    colored left accent, a small bold label chip, and a short executive
    sentence. Used to put the freed-up space below KPI rows / chart rows /
    card rows to use with real content instead of empty whitespace."""
    add_rect(slide, x, y, w, h, fill_hex=(bg or P["bg_alt"]), line_hex=P["border"], line_pt=0.75, rounded=True, radius=0.10)
    add_rect(slide, x, y, Inches(0.07), h, fill_hex=color)
    pad = Inches(0.22)
    add_text(slide, x + pad, y + Inches(0.13), Inches(2.4), Inches(0.24), label.upper(),
              size=9, color=color, bold=True, spacing=30)
    add_text(slide, x + pad, y + Inches(0.40), w - 2 * pad, h - Inches(0.5), text,
              size=11.5, color=P["text2"], line_spacing=1.25)


def add_bullet_card(slide, x, y, w, h, title, bullets, color=P["navy"], title_size=12, bullet_size=10):
    add_rect(slide, x, y, w, h, fill_hex=P["card"], line_hex=P["border"], line_pt=0.75, rounded=True, radius=0.06)
    add_rect(slide, x, y, Inches(0.06), h, fill_hex=color)
    pad = Inches(0.18)
    add_text(slide, x + pad, y + Inches(0.12), w - 2 * pad, Inches(0.28), title,
              size=title_size, color=P["text"], bold=True)
    ty = y + Inches(0.46)
    tb = slide.shapes.add_textbox(x + pad, ty, w - 2 * pad, h - Inches(0.46) - Inches(0.1))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.1
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = f"•  {b}"
        run.font.size = Pt(bullet_size)
        run.font.color.rgb = rgb(P["text2"])
        run.font.name = FONT


# --------------------------------------------------------------------------
# Table
# --------------------------------------------------------------------------
def add_table(slide, x, y, w, h, headers, rows, col_widths,
              header_bg=P["navy"], header_fg=P["white"], font_size=10,
              zebra=True, badge_col=None, badge_fn=None, row_h=None, badge_cols=None,
              zero_muted_cols=None):
    """headers: list[str]; rows: list[list[str]]; col_widths: list[float]
    (relative ratios, normalized automatically). badge_col/badge_fn: a single
    column whose cell should be filled per badge_fn(value) -> (bg_hex, fg_hex).
    badge_cols: for multiple differently-colored columns (e.g. one severity
    count column per severity level) - a dict {col_index: fn(value) -> (bg,
    fg)}, checked in addition to badge_col/badge_fn. zero_muted_cols: column
    indices (typically the same ones as badge_cols) where a value of "0"
    should render as plain muted text instead of a colored badge, so a table
    full of severity-count columns doesn't turn into a wall of colored 0s."""
    badge_cols = badge_cols or {}
    zero_muted_cols = set(zero_muted_cols or [])
    n_rows = len(rows) + 1
    n_cols = len(headers)
    graphic_frame = slide.shapes.add_table(n_rows, n_cols, x, y, w, h)
    table = graphic_frame.table

    total_ratio = sum(col_widths)
    for i, ratio in enumerate(col_widths):
        table.columns[i].width = Emu(int(w * (ratio / total_ratio)))

    if row_h:
        for r in range(n_rows):
            table.rows[r].height = row_h

    # strip the default PowerPoint table style banding so our own colors show
    _strip_table_style(graphic_frame)

    for c, htext in enumerate(headers):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(header_bg)
        cell.margin_left = Inches(0.06)
        cell.margin_right = Inches(0.06)
        cell.margin_top = Inches(0.02)
        cell.margin_bottom = Inches(0.02)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = htext
        run.font.size = Pt(font_size - 0.5)
        run.font.bold = True
        run.font.color.rgb = rgb(header_fg)
        run.font.name = FONT

    for r, row in enumerate(rows, start=1):
        zebra_bg = P["bg_alt"] if (zebra and r % 2 == 0) else P["card"]
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.01)
            cell.margin_bottom = Inches(0.01)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            bg = zebra_bg
            fg = P["text2"]
            bold = False
            is_zero = str(val) == "0"
            if c in badge_cols and not (c in zero_muted_cols and is_zero):
                bg_hex, fg_hex = badge_cols[c](val)
                bg, fg, bold = bg_hex, fg_hex, True
            elif c in zero_muted_cols and is_zero:
                fg = P["muted"]
            elif badge_col is not None and c == badge_col and badge_fn is not None:
                bg_hex, fg_hex = badge_fn(val)
                bg = bg_hex
                fg = fg_hex
                bold = True
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(bg)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = "" if val is None else str(val)
            run.font.size = Pt(font_size)
            run.font.color.rgb = rgb(fg)
            run.font.bold = bold
            run.font.name = FONT
    return graphic_frame


def _strip_table_style(graphic_frame):
    tbl = graphic_frame._element.graphic.graphicData.tbl
    tblPr = tbl.find(qn("a:tblPr"))
    if tblPr is not None:
        tblPr.set("firstRow", "0")
        tblPr.set("bandRow", "0")
        for child in list(tblPr):
            tblPr.remove(child)


def paginate(rows, page_size):
    return [rows[i:i + page_size] for i in range(0, len(rows), page_size)]


# --------------------------------------------------------------------------
# Charts
# --------------------------------------------------------------------------
def add_donut_chart(slide, x, y, w, h, categories, values, colors_hex, title=None):
    data = CategoryChartData()
    data.categories = categories
    data.add_series("series", values)
    gframe = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, x, y, w, h, data)
    chart = gframe.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(9)
    chart.legend.font.name = FONT
    chart.has_title = False
    plot = chart.plots[0]
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = "0"
    dl.number_format_is_linked = False
    dl.font.size = Pt(9)
    dl.font.bold = True
    dl.font.color.rgb = rgb(P["text"])
    series = plot.series[0]
    for i, point in enumerate(series.points):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = rgb(colors_hex[i % len(colors_hex)])
        point.format.line.color.rgb = rgb(P["white"])
        point.format.line.width = Pt(1.5)
    return gframe


def add_bar_chart(slide, x, y, w, h, categories, series_name, values, colors_hex=None, bar_color=None,
                   min_scale=0, max_scale=None):
    data = CategoryChartData()
    data.categories = categories
    data.add_series(series_name, values)
    gframe = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, x, y, w, h, data)
    chart = gframe.chart
    chart.has_legend = False
    chart.has_title = False
    plot = chart.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.font.size = Pt(9)
    dl.font.name = FONT
    dl.font.color.rgb = rgb(P["text"])
    series = plot.series[0]
    if colors_hex:
        for i, point in enumerate(series.points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = rgb(colors_hex[i % len(colors_hex)])
            point.format.line.fill.background()
    else:
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = rgb(bar_color or P["accent"])
        series.format.line.fill.background()
    cat_axis = chart.category_axis
    cat_axis.tick_labels.font.size = Pt(9.5)
    cat_axis.tick_labels.font.name = FONT
    cat_axis.format.line.color.rgb = rgb(P["border"])
    val_axis = chart.value_axis
    if min_scale is not None:
        val_axis.minimum_scale = min_scale
    if max_scale is not None:
        val_axis.maximum_scale = max_scale
    val_axis.tick_labels.font.size = Pt(9)
    val_axis.has_major_gridlines = True
    val_axis.major_gridlines.format.line.color.rgb = rgb(P["border"])
    val_axis.major_gridlines.format.line.width = Pt(0.5)
    val_axis.format.line.fill.background()
    return gframe
